#!/usr/bin/env python3
"""Validate and normalize the Villela SDR Excel export for a one-time MySQL import."""
from __future__ import annotations

import json
import sys
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

EXPECTED = {
    "REUNIÕES": 547,
    "RESPONSÁVEIS": 44,
    "SETORES": 3,
    "VÍNCULOS": 49,
    "HISTÓRICO": 1755,
    "ADVOGADOS": 1,
    "BONIFICAÇÕES": 3,
    "INTEGRAÇÕES_API": 1,
    "LOGS_API": 2324,
    "MIGRAÇÕES": 22,
}
ARCHIVE_ONLY_SHEETS = ["CALENDAR_USUÁRIOS", "CALENDAR_ADVOGADOS", "CONVITES_CALENDAR", "GMAIL_ENVIO"]
SECTOR_CODES = {"PROJURIS_2", "CONCILIE", "TRAFEGO_PAGO"}
ROLES = {"user", "sdr", "gestao", "admin", "master"}
MEETING_STATUSES = {"AGENDADA", "REALIZADA", "NO_SHOW", "CANCELADA"}
CLIENT_STATUSES = {"QUENTE", "MÉDIA", "FRIA"}
CALENDAR_STATUSES = {"NOT_CONNECTED", "PENDING", "SYNCED", "ERROR"}

MEETING_FIELDS = [
    ("id", "id"), ("companyName", "companyName"), ("cnpj", "cnpj"),
    ("contactName", "contactName"), ("phone", "phone"), ("contactDate", "contactDate"),
    ("email", "email"), ("sdrId", "sdrId"), ("responsibleId", "responsibleId"),
    ("lawyer_email", "lawyer_email"), ("lawyer_id", "lawyer_id"),
    ("calendar_owner", "calendar_owner"), ("service", "service"), ("closer", "closer"),
    ("client_status", "client_status"), ("date", "date"), ("time", "time"),
    ("status", "status"), ("notes", "notes"), ("result", "result"),
    ("nextStep", "nextStep"), ("contract_value", "contract_value"),
    ("proposal_value", "proposal_value"), ("closed_contract", "closed_contract"),
    ("closed_contract_value", "closed_contract_value"),
    ("contract_installments", "contract_installments"), ("contract_lawyer", "contract_lawyer"),
    ("nextContactDate", "nextContactDate"), ("noShowReason", "noShowReason"),
    ("pgfnDebtOver2m", "pgfnDebtOver2m"), ("pgfnDebtAmount", "pgfnDebtAmount"),
    ("wasContacted", "wasContacted"), ("wantsReschedule", "wantsReschedule"),
    ("cancelReason", "cancelReason"), ("rescheduledDate", "rescheduledDate"),
    ("rescheduledTime", "rescheduledTime"), ("lead_channel", "lead_channel"),
    ("lead_channel_other", "lead_channel_other"), ("ad_platform", "ad_platform"),
    ("campaign_name", "campaign_name"), ("campaign_id", "campaign_id"),
    ("ad_set_name", "ad_set_name"), ("ad_set_id", "ad_set_id"),
    ("ad_name", "ad_name"), ("ad_id", "ad_id"), ("contact_origin", "contact_origin"),
    ("contact_origin_other", "contact_origin_other"), ("utm_source", "utm_source"),
    ("utm_medium", "utm_medium"), ("utm_campaign", "utm_campaign"),
    ("utm_content", "utm_content"), ("utm_term", "utm_term"),
    ("lead_entry_date", "lead_entry_date"), ("qualification_status", "qualification_status"),
    ("disqualification_reason", "disqualification_reason"),
    ("disqualification_other", "disqualification_other"),
    ("google_calendar_event_id", "google_calendar_event_id"),
    ("google_meet_link", "google_meet_link"),
    ("google_calendar_status", "google_calendar_status"),
    ("google_sync_error", "google_sync_error"), ("google_invite_sent", "google_invite_sent"),
    ("google_invite_sent_at", "google_invite_sent_at"),
    ("google_calendar_synced_at", "google_calendar_synced_at"),
    ("google_meet_created_at", "google_meet_created_at"),
    ("google_calendar_event_url", "google_calendar_event_url"),
    ("google_calendar_event_html_link", "google_calendar_event_html_link"),
    ("createdAt", "createdAt"), ("updatedAt", "updatedAt"),
]


def serial(value: Any) -> Any:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat(timespec="seconds")
    if isinstance(value, bool):
        return value
    if isinstance(value, float) and value.is_integer():
        return int(value)
    return value


def rows_for(workbook, sheet_name: str) -> list[dict[str, Any]]:
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Missing sheet: {sheet_name}")
    iterator = workbook[sheet_name].iter_rows(values_only=True)
    header = next(iterator, None)
    if not header:
        return []
    headers = [str(value).strip() if value is not None else "" for value in header]
    if len([header for header in headers if header]) != len(set(header for header in headers if header)):
        raise ValueError(f"Duplicate column headers in {sheet_name}")
    result = []
    for raw in iterator:
        if not any(value is not None for value in raw):
            continue
        result.append({headers[index]: serial(value) for index, value in enumerate(raw) if index < len(headers) and headers[index]})
    return result


def require_id(value: Any, label: str) -> int:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"Invalid ID in {label}")
    try:
        number = int(value)
    except (TypeError, ValueError):
        raise ValueError(f"Invalid ID in {label}") from None
    if str(number) != str(value).strip() and not (isinstance(value, float) and value.is_integer()):
        raise ValueError(f"Non-integral ID in {label}")
    if number < 1 or number > 2_147_483_647:
        raise ValueError(f"ID out of MySQL INT range in {label}")
    return number


def unique_ids(records: list[dict[str, Any]], label: str) -> set[int]:
    ids = [require_id(record.get("id"), label) for record in records]
    if len(ids) != len(set(ids)):
        raise ValueError(f"Duplicate IDs in {label}")
    return set(ids)


def as_bool(value: Any, label: str, default: Any = None) -> Any:
    if value is None or value == "":
        return default
    if value is True or value == 1 or str(value).strip().lower() in {"1", "true", "sim", "yes"}:
        return 1
    if value is False or value == 0 or str(value).strip().lower() in {"0", "false", "não", "nao", "no"}:
        return 0
    raise ValueError(f"Invalid boolean in {label}")


def make_record(source: dict[str, Any], fields: list[tuple[str, str]]) -> dict[str, Any]:
    return {target: source.get(column) for target, column in fields}


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: prepare-villela-export.py <input.xlsx> <output.json>")
    input_path, output_path = Path(sys.argv[1]), Path(sys.argv[2])
    if not input_path.is_file():
        raise SystemExit("Workbook not found")
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    try:
        manifest_rows = rows_for(workbook, "LEIA-ME")
        manifest = {str(row.get("Item")): serial(row.get("Valor")) for row in manifest_rows if row.get("Item")}
        if manifest.get("Arquivo") != "VILLELA SDR - exportação completa organizada":
            raise ValueError("Unexpected workbook manifest")
        source = {sheet: rows_for(workbook, sheet) for sheet in [*EXPECTED, *ARCHIVE_ONLY_SHEETS, "LEIA-ME"]}
        for sheet, expected in EXPECTED.items():
            if len(source[sheet]) != expected:
                raise ValueError(f"Unexpected row count in {sheet}: {len(source[sheet])}, expected {expected}")

        sectors = source["SETORES"]
        sector_ids = unique_ids(sectors, "SETORES")
        codes = {str(row.get("code", "")).strip() for row in sectors}
        if codes != SECTOR_CODES:
            raise ValueError("Unexpected sector codes in workbook")
        sector_by_id = {require_id(row["id"], "SETORES"): str(row["code"]).strip() for row in sectors}

        users_source = source["RESPONSÁVEIS"]
        user_ids = unique_ids(users_source, "RESPONSÁVEIS")
        for record in users_source:
            if record.get("role") not in ROLES:
                raise ValueError("Unexpected user role")
            email = str(record.get("email") or "").strip()
            if not email or "@" not in email:
                raise ValueError("A user record has an invalid email")
        user_email_counts: dict[str, list[dict[str, Any]]] = {}
        for record in users_source:
            key = str(record["email"]).strip().lower()
            user_email_counts.setdefault(key, []).append(record)
        for group in user_email_counts.values():
            if len(group) > 1 and sum(bool(as_bool(record.get("active"), "RESPONSÁVEIS.active", 0)) for record in group) != 1:
                raise ValueError("Duplicate user email is ambiguous for active login association")

        meetings_source = source["REUNIÕES"]
        meeting_ids = unique_ids(meetings_source, "REUNIÕES")
        for record in meetings_source:
            if str(record.get("status")) not in MEETING_STATUSES:
                raise ValueError("Unexpected meeting status")
            if record.get("client_status") not in CLIENT_STATUSES:
                raise ValueError("Unexpected client status")
            if record.get("google_calendar_status") not in CALENDAR_STATUSES:
                raise ValueError("Unexpected Google Calendar status")
            if require_id(record.get("sdrId"), "REUNIÕES.sdrId") not in user_ids:
                raise ValueError("A meeting references an unknown SDR")
            if require_id(record.get("responsibleId"), "REUNIÕES.responsibleId") not in user_ids:
                raise ValueError("A meeting references an unknown responsible user")
            if str(record.get("sectorCode")) not in SECTOR_CODES:
                raise ValueError("A meeting references an unknown sector code")
            for flag in ("wasContacted", "wantsReschedule", "pgfnDebtOver2m", "closed_contract", "google_invite_sent"):
                record[flag] = as_bool(record.get(flag), f"REUNIÕES.{flag}")
            record["sdrId"] = require_id(record["sdrId"], "REUNIÕES.sdrId")
            record["responsibleId"] = require_id(record["responsibleId"], "REUNIÕES.responsibleId")
            if not record.get("service"):
                record["service"] = ""

        links = source["VÍNCULOS"]
        link_keys = set()
        for record in links:
            uid = require_id(record.get("userId"), "VÍNCULOS.userId")
            if uid not in user_ids or str(record.get("sectorCode")) not in SECTOR_CODES:
                raise ValueError("A user-sector link references an unknown entity")
            key = (uid, str(record["sectorCode"]))
            if key in link_keys:
                raise ValueError("Duplicate user-sector link")
            link_keys.add(key)
            record["userId"] = uid

        history = source["HISTÓRICO"]
        history_ids = unique_ids(history, "HISTÓRICO")
        for record in history:
            if require_id(record.get("meetingId"), "HISTÓRICO.meetingId") not in meeting_ids:
                raise ValueError("History references an unknown meeting")
            if require_id(record.get("userId"), "HISTÓRICO.userId") not in user_ids:
                raise ValueError("History references an unknown user")
            if str(record.get("sectorCode")) not in SECTOR_CODES:
                raise ValueError("History references an unknown sector code")

        lawyers = source["ADVOGADOS"]
        lawyer_ids = unique_ids(lawyers, "ADVOGADOS")
        for record in meetings_source:
            lawyer_id = record.get("lawyer_id")
            if lawyer_id not in (None, "") and require_id(lawyer_id, "REUNIÕES.lawyer_id") not in lawyer_ids:
                raise ValueError("A meeting references an unknown lawyer")

        integrations = source["INTEGRAÇÕES_API"]
        integration_ids = unique_ids(integrations, "INTEGRAÇÕES_API")
        for record in integrations:
            if record.get("sectorCode") != "CONCILIE":
                raise ValueError("Unexpected API integration sector")
            record["active"] = 0  # Its secret hash was intentionally omitted from the export.

        api_logs = source["LOGS_API"]
        unique_ids(api_logs, "LOGS_API")
        for record in api_logs:
            if require_id(record.get("integrationId"), "LOGS_API.integrationId") not in integration_ids:
                raise ValueError("An API log references an unknown integration")

        bonuses = source["BONIFICAÇÕES"]
        unique_ids(bonuses, "BONIFICAÇÕES")
        for record in bonuses:
            source_sector_id = require_id(record.get("sectorId"), "BONIFICAÇÕES.sectorId")
            if source_sector_id not in sector_by_id:
                raise ValueError("A bonus record references an unknown sector")
            record["sectorCode"] = sector_by_id[source_sector_id]

        users = []
        for record in users_source:
            active = as_bool(record.get("active"), "RESPONSÁVEIS.active", 0)
            users.append({
                "id": require_id(record["id"], "RESPONSÁVEIS.id"),
                "openId": f"villela_import_{require_id(record['id'], 'RESPONSÁVEIS.id')}",
                "name": record.get("name"),
                "email": str(record["email"]).strip().lower(),
                "loginMethod": record.get("loginMethod") or "imported",
                "role": record["role"],
                "active": active,
                "createdAt": record.get("createdAt"),
                "updatedAt": record.get("updatedAt"),
                "lastSignedIn": record.get("lastSignedIn") or record.get("createdAt"),
            })

        meetings = []
        for record in meetings_source:
            mapped = make_record(record, MEETING_FIELDS)
            mapped["id"] = require_id(record["id"], "REUNIÕES.id")
            mapped["sectorCode"] = str(record["sectorCode"])
            mapped["sdrId"] = require_id(record["sdrId"], "REUNIÕES.sdrId")
            mapped["responsibleId"] = require_id(record["responsibleId"], "REUNIÕES.responsibleId")
            meetings.append(mapped)

        meeting_history = []
        for record in history:
            meeting_history.append({
                "id": require_id(record["id"], "HISTÓRICO.id"),
                "meetingId": require_id(record["meetingId"], "HISTÓRICO.meetingId"),
                "userId": require_id(record["userId"], "HISTÓRICO.userId"),
                "sectorCode": str(record["sectorCode"]),
                "action": record.get("action"), "oldValue": record.get("oldValue"),
                "newValue": record.get("newValue"), "createdAt": record.get("createdAt"),
            })

        sector_records = [{
            "id": require_id(row["id"], "SETORES.id"), "code": str(row["code"]).strip(),
            "name": row.get("name"), "active": as_bool(row.get("active"), "SETORES.active", 1),
            "createdAt": row.get("createdAt"),
        } for row in sectors]
        lawyers_out = [{
            "id": require_id(row["id"], "ADVOGADOS.id"), "name": row.get("name"),
            "email": str(row.get("email") or "").strip().lower(),
            "active": as_bool(row.get("active"), "ADVOGADOS.active", 1),
            "createdAt": row.get("createdAt"), "updatedAt": row.get("updatedAt"),
        } for row in lawyers]
        bonuses_out = [{
            "id": require_id(row["id"], "BONIFICAÇÕES.id"), "monthKey": row.get("monthKey"),
            "sectorCode": str(row["sectorCode"]), "status": row.get("status"),
            "closedAt": row.get("closedAt"), "snapshot": row.get("snapshot"),
            "rankingSnapshot": row.get("rankingSnapshot"), "createdAt": row.get("createdAt"),
            "updatedAt": row.get("updatedAt"),
        } for row in bonuses]
        integrations_out = [{
            "id": require_id(row["id"], "INTEGRAÇÕES_API.id"), "name": row.get("name"),
            "keyPrefix": row.get("keyPrefix"), "sectorCode": row.get("sectorCode"),
            "permissions": row.get("permissions"), "active": 0,
            "createdAt": row.get("createdAt"), "revokedAt": row.get("revokedAt"),
            "lastUsedAt": row.get("lastUsedAt"), "secretHash": None,
        } for row in integrations]
        logs_out = [{
            "id": require_id(row["id"], "LOGS_API.id"), "integrationId": require_id(row["integrationId"], "LOGS_API.integrationId"),
            "method": row.get("method"), "endpoint": row.get("endpoint"), "statusCode": row.get("statusCode"),
            "ip": row.get("ip"), "userAgent": row.get("userAgent"), "createdAt": row.get("createdAt"),
        } for row in api_logs]

        archive_sheets = {
            "CALENDAR_USUÁRIOS": "calendar_user",
            "CALENDAR_ADVOGADOS": "calendar_lawyer",
            "CONVITES_CALENDAR": "calendar_invite",
            "GMAIL_ENVIO": "gmail_sender",
            "MIGRAÇÕES": "source_migration",
        }
        archived = [{"sourceType": "export_manifest", "sourceId": 1, "payload": manifest}]
        for sheet, source_type in archive_sheets.items():
            for row in source[sheet]:
                archived.append({
                    "sourceType": source_type,
                    "sourceId": require_id(row.get("id"), sheet),
                    "payload": row,
                })
        for row in source["LEIA-ME"]:
            if row.get("Item") and row.get("Item") != "Valor":
                archived.append({"sourceType": "export_note", "sourceId": len(archived) + 1, "payload": row})

        output = {
            "sourceFile": input_path.name,
            "sourceCounts": {sheet: len(source[sheet]) for sheet in source},
            "manifest": manifest,
            "sectors": sector_records,
            "users": users,
            "userSectors": [{"userId": record["userId"], "sectorCode": record["sectorCode"], "createdAt": record.get("createdAt")} for record in links],
            "meetings": meetings,
            "meetingHistory": meeting_history,
            "lawyers": lawyers_out,
            "bonusCompetencies": bonuses_out,
            "apiIntegrations": integrations_out,
            "apiIntegrationLogs": logs_out,
            "archivedMetadata": archived,
        }
        output_path.parent.mkdir(parents=True, exist_ok=True)
        with output_path.open("w", encoding="utf-8") as out:
            json.dump(output, out, ensure_ascii=False, separators=(",", ":"), default=str)
        output_path.chmod(0o600)
        print(json.dumps({"valid": True, "sourceCounts": output["sourceCounts"], "targetCounts": {
            "users": len(users), "sectors": len(sector_records), "userSectors": len(links),
            "meetings": len(meetings), "meetingHistory": len(meeting_history), "lawyers": len(lawyers_out),
            "bonusCompetencies": len(bonuses_out), "apiIntegrations": len(integrations_out),
            "apiIntegrationLogs": len(logs_out), "archivedMetadata": len(archived),
        }}))
    finally:
        workbook.close()


if __name__ == "__main__":
    main()
